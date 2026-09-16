# V1.4
"""
Backtest the bot's strategies against your OWN recorded live-tick JSONL data
(written by data_logger.py into ./data/), instead of fetching from the Kalshi
API. This never calls Kalshi or places any order - it's pure offline replay.
... [rest of the docstring remains the same] ...
"""
from __future__ import annotations

import argparse
import bisect
import datetime as dt
import glob
import itertools
import json
import logging
import os
import re
import statistics
import sys
import tempfile
import time
from dataclasses import dataclass, field
from typing import Optional

import yaml

from state import StateStore, PendingBet
from strategy import (
    decide_side,
    decide_price_trend_side,
    decide_spot_lean_side,
    decide_late_fade_side,
    both_sides_too_expensive,
    compute_recovery_size,
    compute_smart_hedge_count,
    compute_take_profit_profit,
    check_momentum_filter,
)
from bot import contracts_for_stake, score_pending_bets, _session_side_totals

log = logging.getLogger("tick_backtest")

FILENAME_RE = re.compile(r"^(?P<stem>.+)_(?P<date>\d{8})_(?P<time>\d{4})\.jsonl$")

# ---------------------------------------------------------------------------
# Data loading
# ---------------------------------------------------------------------------

@dataclass
class Tick:
    t: dt.datetime
    spot: Optional[float]
    strike: Optional[float]
    up_cents: Optional[int]
    down_cents: Optional[int]

@dataclass
class WindowData:
    open_time: dt.datetime
    close_time: dt.datetime
    strike: Optional[float] = None
    ticks: list = field(default_factory=list)
    path: str = ""

def parse_filename_open_time(path: str) -> Optional[dt.datetime]:
    m = FILENAME_RE.match(os.path.basename(path))
    if not m:
        return None
    try:
        return dt.datetime.strptime(m.group("date") + m.group("time"), "%Y%m%d%H%M")
    except ValueError:
        return None

def load_window_file(path: str) -> Optional[WindowData]:
    ticks = []
    with open(path, "r", encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            try:
                rec = json.loads(line)
            except json.JSONDecodeError:
                continue
            t_raw = rec.get("time")
            if not t_raw:
                continue
            try:
                t = dt.datetime.fromisoformat(t_raw)
            except ValueError:
                continue
            ticks.append(Tick(
                t=t,
                spot=rec.get("btc_spot_cfbenchmarks"),
                strike=rec.get("kalshi_strike_usd"),
                up_cents=rec.get("kalshi_up_price_cents"),
                down_cents=rec.get("kalshi_down_price_cents"),
            ))
    if not ticks:
        return None
    ticks.sort(key=lambda tk: tk.t)

    open_time = parse_filename_open_time(path)
    if open_time is None:
        first = ticks[0].t
        floor_minute = (first.minute // 15) * 15
        open_time = first.replace(minute=floor_minute, second=0, microsecond=0)
    close_time = open_time + dt.timedelta(minutes=15)

    strikes = [tk.strike for tk in ticks if tk.strike]
    strike = statistics.mode(strikes) if strikes else None

    return WindowData(open_time=open_time, close_time=close_time, strike=strike, ticks=ticks, path=path)

def load_all_windows(data_dir: str, pattern: str = "*.jsonl") -> list:
    paths = sorted(glob.glob(os.path.join(data_dir, pattern)))
    windows = []
    for p in paths:
        w = load_window_file(p)
        if w is not None:
            windows.append(w)
        else:
            log.debug("Skipping empty/unparseable file: %s", p)
    windows.sort(key=lambda w: w.open_time)
    return windows

def contiguous_segments(windows: list, tolerance_min: float = 1.0) -> list:
    if not windows:
        return []
    segments = [[windows[0]]]
    for prev, cur in zip(windows, windows[1:]):
        gap_min = (cur.open_time - prev.close_time).total_seconds() / 60.0
        if abs(gap_min) <= tolerance_min:
            segments[-1].append(cur)
        else:
            segments.append([cur])
    return segments

def determine_result(window: WindowData, next_window: Optional[WindowData]) -> Optional[str]:
    if window.strike is not None and next_window is not None and next_window.strike is not None:
        if next_window.strike > window.strike:
            return "yes"
        if next_window.strike < window.strike:
            return "no"
        return None
    if window.strike is not None:
        for tk in reversed(window.ticks):
            if tk.spot is not None:
                if tk.spot > window.strike:
                    return "yes"
                if tk.spot < window.strike:
                    return "no"
                return None
    return None

def build_spot_index(windows: list):
    series = []
    for w in windows:
        for tk in w.ticks:
            if tk.spot is not None:
                series.append((tk.t, tk.spot))
    series.sort(key=lambda p: p[0])
    times = [t for t, _ in series]
    prices = [p for _, p in series]
    return times, prices

def momentum_history_slice(times: list, prices: list, now_dt: dt.datetime, lookback_sec: float) -> list:
    lo_time = now_dt - dt.timedelta(seconds=lookback_sec)
    lo = bisect.bisect_left(times, lo_time)
    hi = bisect.bisect_right(times, now_dt)
    return [(times[j].timestamp(), prices[j]) for j in range(lo, hi)]

# ---------------------------------------------------------------------------
# Shared bet-placement gate
# ---------------------------------------------------------------------------

def _try_place(tk: Tick, side: str, window: WindowData, ticker: str, cfg: dict, store: StateStore, spot_index):
    price = tk.up_cents if side == "yes" else tk.down_cents
    if price is None:
        return None

    max_price = cfg["strategy"]["max_price_cents"]
    min_price = cfg["strategy"].get("min_price_cents", 1)
    sizing_mode = cfg["sizing"]["mode"]
    fee = cfg["sizing"].get("fee_per_contract_cents", 0)

    mf_cfg = cfg["strategy"].get("momentum_filter", {})
    if mf_cfg.get("enabled", False):
        mf_lookback = mf_cfg.get("lookback_sec", 30)
        hist = momentum_history_slice(spot_index[0], spot_index[1], tk.t, mf_lookback)
        allowed, _direction = check_momentum_filter(hist, mf_lookback, side, now=tk.t.timestamp())
        if not allowed:
            return None

    recovering = sizing_mode == "recovery" and store.state.cumulative_loss_cents > 0
    if recovering:
        rc = cfg.get("recovery", {})
        r_min_profit = rc.get("min_profit_cents", 1)
        r_max_contracts = rc.get("max_contracts", 50)
        r_max_price = rc.get("max_price_cents", max_price)
        r_min_price = rc.get("min_price_cents", min_price)
        if price > r_max_price or price < r_min_price:
            return None
        count = compute_recovery_size(store.state.cumulative_loss_cents, price, fee, r_min_profit, r_max_contracts)
        if count is None:
            return None
    else:
        if not (min_price <= price <= max_price):
            return None
        count = contracts_for_stake(store.state.current_stake, sizing_mode, price)

    store.add_pending_bet(PendingBet(
        window_close_iso=window.close_time.isoformat(), ticker=ticker,
        side=side, stake=count, price_cents=price, order_id=None,
    ))
    return {"price": price, "count": count, "tick": tk, "side": side}

def _scan_entry_fixed_side(window: WindowData, side: str, cfg: dict, store: StateStore, ticker: str, spot_index):
    entry_start = cfg["strategy"]["entry_start_min"]
    entry_end = cfg["strategy"]["entry_end_min"]
    for tk in window.ticks:
        elapsed_min = (tk.t - window.open_time).total_seconds() / 60.0
        if elapsed_min < entry_start:
            continue
        if elapsed_min > entry_end:
            break
        placed = _try_place(tk, side, window, ticker, cfg, store, spot_index)
        if placed:
            return placed
    return None

def _scan_spot_lean_entry(window: WindowData, target: float, cfg: dict, store: StateStore, ticker: str, spot_index):
    sl_cfg = cfg["strategy"]["spot_lean"]
    threshold_pct = sl_cfg.get("threshold_pct", 0.0)
    entry_start = cfg["strategy"]["entry_start_min"]
    entry_end = cfg["strategy"]["entry_end_min"]
    for tk in window.ticks:
        elapsed_min = (tk.t - window.open_time).total_seconds() / 60.0
        if elapsed_min < entry_start:
            continue
        if elapsed_min > entry_end:
            break
        side, _gap = decide_spot_lean_side(tk.spot, target, threshold_pct)
        if side is None:
            continue
        placed = _try_place(tk, side, window, ticker, cfg, store, spot_index)
        if placed:
            return placed
    return None

def _scan_late_fade_entry(window: WindowData, target: float, cfg: dict, store: StateStore, ticker: str, spot_index):
    """
    late_fade: same tick-by-tick scan as spot_lean, but (1) decides the side
    via decide_late_fade_side using a min/max threshold range (bets a reversal 
    back toward `target`), and (2) applies bot.py's session-skip gate.
    """
    lf_cfg = cfg["strategy"].get("late_fade", {})
    min_threshold_pct = lf_cfg.get("min_threshold_pct", 0.0)
    max_threshold_pct = lf_cfg.get("max_threshold_pct", 0.0)
    
    entry_start = cfg["strategy"]["entry_start_min"]
    entry_end = cfg["strategy"]["entry_end_min"]
    max_price = cfg["strategy"]["max_price_cents"]

    gate_checked = False
    for tk in window.ticks:
        elapsed_min = (tk.t - window.open_time).total_seconds() / 60.0
        if elapsed_min < entry_start:
            continue
        if elapsed_min > entry_end:
            break
            
        if not gate_checked:
            gate_checked = True
            if both_sides_too_expensive(tk.up_cents, tk.down_cents, max_price):
                return None  # session-skip gate: neither side cheap enough right at entry_start
                
        side, _gap = decide_late_fade_side(tk.spot, target, min_threshold_pct, max_threshold_pct)
        if side is None:
            continue
            
        placed = _try_place(tk, side, window, ticker, cfg, store, spot_index)
        if placed:
            return placed
            
    return None

def _scan_spot_lean_hedges_and_take_profit(
    window: WindowData, target: float, main_placed: dict, cfg: dict, store: StateStore, ticker: str,
    observer=None,
) -> dict:
    """
    FIXED: Now reads hedge/take_profit from top-level strategy config, 
    matching bot.py's monitor_hedge logic, rather than hardcoding "spot_lean".
    """
    # Hedge and take-profit are now top-level strategy configs, independent of mode
    hedge_cfg = cfg["strategy"].get("hedge", {})
    tp_cfg = cfg["strategy"].get("take_profit", {})

    hedge_enabled = hedge_cfg.get("enabled", False)
    tp_enabled = tp_cfg.get("enabled", False)
    if not hedge_enabled and not tp_enabled:
        return {"hedges_placed": 0, "take_profit_placed": False}

    if hedge_enabled and hedge_cfg.get("fresh_start_only", True) and main_placed["count"] != 1:
        hedge_enabled = False
    if tp_enabled and tp_cfg.get("fresh_start_only", True) and main_placed["count"] != 1:
        tp_enabled = False
    if not hedge_enabled and not tp_enabled:
        return {"hedges_placed": 0, "take_profit_placed": False}

    hedge_threshold = hedge_cfg.get("threshold_pct", 0.0)
    max_hedges = hedge_cfg.get("max_hedges_per_window", 1)
    min_price = cfg["strategy"].get("min_price_cents", 1)
    max_price = cfg["strategy"]["max_price_cents"]
    smart_sizing = hedge_cfg.get("smart_sizing", False)
    smart_min_profit = hedge_cfg.get("min_profit_cents", 1)
    smart_max_contracts = hedge_cfg.get("max_contracts", 50)
    fee = cfg["sizing"].get("fee_per_contract_cents", 0)

    tp_threshold_pct = tp_cfg.get("threshold_pct", 0.0)
    tp_min_profit = tp_cfg.get("min_profit_cents", 1)
    tp_min_price = tp_cfg.get("min_price_cents", 1)
    tp_max_price = tp_cfg.get("max_price_cents", max_price)
    tp_time_start = tp_cfg.get("time_start_min", 0)
    tp_time_end = tp_cfg.get("time_end_min", 15)
    initial_side = main_placed["side"]

    current_side = main_placed["side"]
    hedges_placed = 0

    for tk in window.ticks:
        if tk.t <= main_placed["tick"].t:
            continue
        if not hedge_enabled and not tp_enabled:
            break

        if tp_enabled:
            minutes_into_window = (tk.t - window.open_time).total_seconds() / 60.0
            if tp_time_start <= minutes_into_window <= tp_time_end:
                tp_side, _tp_gap = decide_spot_lean_side(tk.spot, target, tp_threshold_pct)
                if tp_side == initial_side:
                    opposite_side = "no" if initial_side == "yes" else "yes"
                    tp_price = tk.up_cents if opposite_side == "yes" else tk.down_cents
                    if tp_price is not None and tp_min_price <= tp_price <= tp_max_price:
                        totals = _session_side_totals(store, ticker, fee)
                        main_count_now, main_cost_now = totals[initial_side]
                        if main_count_now > 0:
                            guaranteed_profit = compute_take_profit_profit(
                                main_placed["count"], main_cost_now, tp_price, fee,
                            )
                            if guaranteed_profit >= tp_min_profit:
                                store.add_pending_bet(PendingBet(
                                    window_close_iso=window.close_time.isoformat(), ticker=ticker,
                                    side=opposite_side, stake=main_placed["count"], price_cents=tp_price,
                                    order_id=None,
                                ))
                                if observer:
                                    observer.order_placed(
                                        window, tk.t, ticker, opposite_side, tp_price,
                                        main_placed["count"], "take_profit",
                                    )
                                return {"hedges_placed": hedges_placed, "take_profit_placed": True}

        if not hedge_enabled or hedges_placed >= max_hedges:
            continue

        side, _gap = decide_spot_lean_side(tk.spot, target, hedge_threshold)
        if side is None or side == current_side:
            continue
        price = tk.up_cents if side == "yes" else tk.down_cents
        if price is None or not (min_price <= price <= max_price):
            continue

        if smart_sizing:
            totals = _session_side_totals(store, ticker, fee)
            this_count, this_cost = totals[side]
            other_side = "no" if side == "yes" else "yes"
            _other_count, other_cost = totals[other_side]
            hedge_count, _guaranteed = compute_smart_hedge_count(
                this_count, this_cost, other_cost, price, fee, smart_min_profit, smart_max_contracts,
            )
            if hedge_count is None:
                current_side = side
                continue
            count = hedge_count
        else:
            count = main_placed["count"]

        store.add_pending_bet(PendingBet(
            window_close_iso=window.close_time.isoformat(), ticker=ticker,
            side=side, stake=count, price_cents=price, order_id=None,
        ))
        if observer:
            observer.order_placed(window, tk.t, ticker, side, price, count, "hedge")
        hedges_placed += 1
        current_side = side

    return {"hedges_placed": hedges_placed, "take_profit_placed": False}

# ---------------------------------------------------------------------------
# Results
# ---------------------------------------------------------------------------

@dataclass
class TickBacktestResult:
    label: str
    bets: int = 0
    wins: int = 0
    losses: int = 0
    skipped: int = 0
    max_loss_streak: int = 0
    max_drawdown_cents: int = 0
    realized_pnl_cents: int = 0
    total_price_paid_cents: int = 0
    total_contracts: int = 0
    hedges_placed: int = 0
    take_profits_placed: int = 0

    @property
    def win_rate(self):
        return (self.wins / self.bets * 100.0) if self.bets else None

    @property
    def avg_price(self):
        return (self.total_price_paid_cents / self.total_contracts) if self.total_contracts else None

def _score_and_record(stats: TickBacktestResult, store: StateStore, cfg: dict, ticker: str,
                       window: WindowData, result: Optional[str], price_cost_cents: int, contracts: int,
                       loss_streak: int, observer=None) -> int:
    if result is None:
        store.state.pending_bets = [b for b in store.state.pending_bets if b["ticker"] != ticker]
        stats.skipped += 1
        return loss_streak

    before_pnl = store.state.total_pnl_cents
    before_wins = store.state.total_wins
    before_losses = store.state.total_losses
    score_pending_bets(store, cfg, ticker, result)
    pnl_delta = store.state.total_pnl_cents - before_pnl

    if observer is not None:
        observer.settled(window, ticker, contracts, pnl_delta, store.state.total_pnl_cents)

    stats.bets += (store.state.total_wins - before_wins) + (store.state.total_losses - before_losses)
    stats.wins += store.state.total_wins - before_wins
    stats.losses += store.state.total_losses - before_losses
    stats.realized_pnl_cents += pnl_delta
    stats.total_price_paid_cents += price_cost_cents
    stats.total_contracts += contracts

    if pnl_delta >= 0:
        return 0
    loss_streak += 1
    stats.max_loss_streak = max(stats.max_loss_streak, loss_streak)
    return loss_streak

# ---------------------------------------------------------------------------
# Per-strategy segment simulators
# ---------------------------------------------------------------------------

def simulate_spot_lean_segment(segment, results, spot_index, cfg, store, observer=None) -> TickBacktestResult:
    stats = TickBacktestResult(label="spot_lean")
    loss_streak = 0
    for i, window in enumerate(segment):
        if observer:
            observer.new_window(window)
        if window.strike is None:
            stats.skipped += 1
            continue
        ticker = f"SIM_{window.open_time.isoformat()}"
        main = _scan_spot_lean_entry(window, window.strike, cfg, store, ticker, spot_index)
        if main is None:
            stats.skipped += 1
            continue
        if observer:
            observer.order_placed(window, main["tick"].t, ticker, main["side"], main["price"], main["count"], "main")
        outcome = _scan_spot_lean_hedges_and_take_profit(
            window, window.strike, main, cfg, store, ticker, observer=observer,
        )
        stats.hedges_placed += outcome["hedges_placed"]
        if outcome["take_profit_placed"]:
            stats.take_profits_placed += 1

        placed_bets = [b for b in store.state.pending_bets if b["ticker"] == ticker]
        cost = sum(b["stake"] * b["price_cents"] for b in placed_bets)
        contracts = sum(b["stake"] for b in placed_bets)
        if observer:
            observer.before_score(window)
        loss_streak = _score_and_record(
            stats, store, cfg, ticker, window, results[i], cost, contracts, loss_streak, observer=observer,
        )
    return stats

def simulate_late_fade_segment(segment, results, spot_index, cfg, store, observer=None) -> TickBacktestResult:
    stats = TickBacktestResult(label="late_fade")
    loss_streak = 0
    for i, window in enumerate(segment):
        if observer:
            observer.new_window(window)
        if window.strike is None:
            stats.skipped += 1
            continue
        ticker = f"SIM_{window.open_time.isoformat()}"
        main = _scan_late_fade_entry(window, window.strike, cfg, store, ticker, spot_index)
        if main is None:
            stats.skipped += 1
            continue
        if observer:
            observer.order_placed(window, main["tick"].t, ticker, main["side"], main["price"], main["count"], "main")
        outcome = _scan_spot_lean_hedges_and_take_profit(
            window, window.strike, main, cfg, store, ticker, observer=observer,
        )
        stats.hedges_placed += outcome["hedges_placed"]
        if outcome["take_profit_placed"]:
            stats.take_profits_placed += 1

        placed_bets = [b for b in store.state.pending_bets if b["ticker"] == ticker]
        cost = sum(b["stake"] * b["price_cents"] for b in placed_bets)
        contracts = sum(b["stake"] for b in placed_bets)
        if observer:
            observer.before_score(window)
        loss_streak = _score_and_record(
            stats, store, cfg, ticker, window, results[i], cost, contracts, loss_streak, observer=observer,
        )
    return stats

def simulate_fixed_mode_segment(segment, results, spot_index, cfg, store, mode: str, observer=None) -> TickBacktestResult:
    stats = TickBacktestResult(label=mode)
    loss_streak = 0
    if observer and segment:
        observer.new_window(segment[0])
    for i in range(1, len(segment)):
        prev_result = results[i - 1]
        window = segment[i]
        if observer:
            observer.new_window(window)
        if prev_result is None or window.strike is None:
            stats.skipped += 1
            continue
        side = decide_side(prev_result, mode)
        ticker = f"SIM_{window.open_time.isoformat()}"
        placed = _scan_entry_fixed_side(window, side, cfg, store, ticker, spot_index)
        if not placed:
            stats.skipped += 1
            continue
        if observer:
            observer.order_placed(window, placed["tick"].t, ticker, placed["side"], placed["price"], placed["count"], "main")
            observer.before_score(window)
        loss_streak = _score_and_record(
            stats, store, cfg, ticker, window, results[i],
            placed["price"] * placed["count"], placed["count"], loss_streak, observer=observer,
        )
    return stats

def simulate_adaptive_segment(segment, results, spot_index, cfg, store, default_mode: str, observer=None) -> TickBacktestResult:
    stats = TickBacktestResult(label=f"adaptive(default={default_mode})")
    loss_streak = 0
    if observer and segment:
        observer.new_window(segment[0])
    for i in range(1, len(segment)):
        prev_result = results[i - 1]
        window = segment[i]
        if observer:
            observer.new_window(window)
        if prev_result is None or window.strike is None:
            stats.skipped += 1
            continue
        mode = default_mode if store.state.last_bet_won is None else (
            "momentum" if store.state.last_bet_won else "reversal"
        )
        side = decide_side(prev_result, mode)
        ticker = f"SIM_{window.open_time.isoformat()}"
        placed = _scan_entry_fixed_side(window, side, cfg, store, ticker, spot_index)
        if not placed:
            stats.skipped += 1
            continue
        if observer:
            observer.order_placed(window, placed["tick"].t, ticker, placed["side"], placed["price"], placed["count"], "main")
            observer.before_score(window)
        loss_streak = _score_and_record(
            stats, store, cfg, ticker, window, results[i],
            placed["price"] * placed["count"], placed["count"], loss_streak, observer=observer,
        )
    return stats

def simulate_price_trend_segment(segment, results, spot_index, cfg, store, observer=None) -> TickBacktestResult:
    pt_cfg = cfg["strategy"].get("price_trend", {})
    lookback = pt_cfg.get("lookback_cycles", 6)
    threshold_pct = pt_cfg.get("threshold_pct", 0.15)
    stats = TickBacktestResult(label=f"price_trend(lookback={lookback}, threshold={threshold_pct})")
    loss_streak = 0
    if observer:
        for w in segment[:lookback]:
            observer.new_window(w)
    for i in range(lookback, len(segment)):
        window_slice = segment[i - lookback: i + 1]
        series = [(None, w.strike) for w in window_slice]
        side, _pct = decide_price_trend_side(series, threshold_pct)
        window = segment[i]
        if observer:
            observer.new_window(window)
        if side is None:
            stats.skipped += 1
            continue
        ticker = f"SIM_{window.open_time.isoformat()}"
        placed = _scan_entry_fixed_side(window, side, cfg, store, ticker, spot_index)
        if not placed:
            stats.skipped += 1
            continue
        if observer:
            observer.order_placed(window, placed["tick"].t, ticker, placed["side"], placed["price"], placed["count"], "main")
            observer.before_score(window)
        loss_streak = _score_and_record(
            stats, store, cfg, ticker, window, results[i],
            placed["price"] * placed["count"], placed["count"], loss_streak, observer=observer,
        )
    return stats

def run_strategy_over_segments(
    strategy: str, segments: list, results_by_segment: list, spot_index, cfg: dict, observer=None,
) -> TickBacktestResult:
    agg = TickBacktestResult(label=strategy)
    for seg, results in zip(segments, results_by_segment):
        if len(seg) < 2:
            continue
        fd, state_path = tempfile.mkstemp(suffix=".json")
        os.close(fd)
        try:
            store = StateStore(state_path, base_stake=cfg["sizing"]["base_size"])
            if strategy == "spot_lean":
                r = simulate_spot_lean_segment(seg, results, spot_index, cfg, store, observer=observer)
            elif strategy == "late_fade":
                r = simulate_late_fade_segment(seg, results, spot_index, cfg, store, observer=observer)
            elif strategy == "price_trend":
                r = simulate_price_trend_segment(seg, results, spot_index, cfg, store, observer=observer)
            elif strategy == "adaptive":
                default_mode = cfg["strategy"].get("adaptive_default_mode", "momentum")
                r = simulate_adaptive_segment(seg, results, spot_index, cfg, store, default_mode, observer=observer)
            elif strategy in ("momentum", "reversal"):
                r = simulate_fixed_mode_segment(seg, results, spot_index, cfg, store, strategy, observer=observer)
            else:
                raise ValueError(f"Unknown strategy: {strategy}")

            if observer:
                observer.segment_finished(store.state)

            agg.bets += r.bets
            agg.wins += r.wins
            agg.losses += r.losses
            agg.skipped += r.skipped
            agg.realized_pnl_cents += r.realized_pnl_cents
            agg.total_price_paid_cents += r.total_price_paid_cents
            agg.total_contracts += r.total_contracts
            agg.hedges_placed += r.hedges_placed
            agg.take_profits_placed += r.take_profits_placed
            agg.max_loss_streak = max(agg.max_loss_streak, r.max_loss_streak)
            agg.max_drawdown_cents = max(agg.max_drawdown_cents, store.state.max_drawdown_cents)
        finally:
            try:
                os.unlink(state_path)
            except OSError:
                pass
    return agg

# ---------------------------------------------------------------------------
# Config grid handling
# ---------------------------------------------------------------------------

def load_yaml(path: str) -> dict:
    with open(path, "r", encoding="utf-8") as f:
        return yaml.safe_load(f) or {}

def deep_copy(d):
    return json.loads(json.dumps(d))

def flatten_grid(node, prefix=()):
    leaves = []
    if isinstance(node, dict):
        for k, v in node.items():
            leaves.extend(flatten_grid(v, prefix + (k,)))
    elif isinstance(node, list):
        leaves.append((prefix, node))
    else:
        leaves.append((prefix, [node]))
    return leaves

def set_path(d: dict, path: tuple, value):
    cur = d
    for k in path[:-1]:
        cur = cur.setdefault(k, {})
    cur[path[-1]] = value

def _norm_path(path: tuple) -> tuple:
    return path[1:] if path and path[0] == "strategy" else path

def iter_grid_combos(base_cfg: dict, *grids):
    leaves = []
    for g in grids:
        leaves.extend(flatten_grid(g or {}))
    if not leaves:
        yield {}, deep_copy(base_cfg)
        return
    paths = [p for p, _ in leaves]
    value_lists = [v for _, v in leaves]
    for combo in itertools.product(*value_lists):
        cfg = deep_copy(base_cfg)
        params = {}
        for path, value in zip(paths, combo):
            set_path(cfg, path, value)
            params[_norm_path(path)] = value
        yield params, cfg

def build_label(params: dict, label_fields: Optional[list] = None) -> str:
    if not params:
        return "(default config)"
    if label_fields is None:
        selected = list(params.items())
    else:
        selected = []
        for field_name in label_fields:
            wanted = _norm_path(tuple(field_name.strip().split(".")))
            for path, value in params.items():
                if path == wanted:
                    selected.append((path, value))
                    break
        if not selected:
            return "(no matching label fields)"
    return ", ".join(f"{'.'.join(path)}={value}" for path, value in selected)

SORT_KEYS = {
    "pnl": lambda r: r.realized_pnl_cents,
    "bets": lambda r: r.bets,
    "wins": lambda r: r.wins,
    "losses": lambda r: r.losses,
    "skipped": lambda r: r.skipped,
    "win_rate": lambda r: r.win_rate if r.win_rate is not None else -1.0,
    "avg_price": lambda r: r.avg_price if r.avg_price is not None else -1.0,
    "max_loss_streak": lambda r: r.max_loss_streak,
    "max_drawdown": lambda r: r.max_drawdown_cents,
    "hedges": lambda r: r.hedges_placed,
    "take_profits": lambda r: r.take_profits_placed,
}

STAT_COLUMNS = [
    "bets", "wins", "losses", "skipped", "win_rate_pct", "avg_price_cents",
    "max_loss_streak", "max_drawdown_usd", "pnl_usd", "hedges_placed", "take_profits_placed",
]

def rank_results(results: list, sort_by: str, sort_order: str) -> list:
    key_fn = SORT_KEYS.get(sort_by, SORT_KEYS["pnl"])
    return sorted(results, key=lambda kv: key_fn(kv[1]), reverse=(sort_order != "asc"))

def result_stat_row(r: "TickBacktestResult") -> dict:
    return {
        "bets": r.bets,
        "wins": r.wins,
        "losses": r.losses,
        "skipped": r.skipped,
        "win_rate_pct": round(r.win_rate, 2) if r.win_rate is not None else "",
        "avg_price_cents": round(r.avg_price, 2) if r.avg_price is not None else "",
        "max_loss_streak": r.max_loss_streak,
        "max_drawdown_usd": round(r.max_drawdown_cents / 100, 2),
        "pnl_usd": round(r.realized_pnl_cents / 100, 2),
        "hedges_placed": r.hedges_placed,
        "take_profits_placed": r.take_profits_placed,
    }

def _varying_param_columns(param_rows: list) -> list:
    order = []
    seen_cols = set()
    values_by_col = {}
    for params in param_rows:
        for path_tuple, value in params.items():
            col = ".".join(path_tuple)
            if col not in seen_cols:
                seen_cols.add(col)
                order.append(col)
            values_by_col.setdefault(col, set())
            try:
                values_by_col[col].add(value)
            except TypeError:
                values_by_col[col].add(str(value))
    return [col for col in order if len(values_by_col[col]) > 1]

def export_results_csv(path: str, all_results: dict, sort_by: str, sort_order: str):
    import csv

    all_params = [params for results in all_results.values() for params, _ in results]
    param_columns = _varying_param_columns(all_params)

    fieldnames = ["strategy"] + param_columns + STAT_COLUMNS
    with open(path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        for strategy, results in all_results.items():
            for params, r in rank_results(results, sort_by, sort_order):
                row = {"strategy": strategy}
                for path_tuple, value in params.items():
                    col = ".".join(path_tuple)
                    if col in param_columns:
                        row[col] = value
                row.update(result_stat_row(r))
                writer.writerow(row)
    log.info("Wrote %s combo(s) across %s strategy/strategies to %s", sum(len(v) for v in all_results.values()), len(all_results), path)

def _drawdown_color_bucket(max_drawdown_usd: float) -> str:
    if max_drawdown_usd < 100.0:
        return "dark_green"
    if max_drawdown_usd < 200.0:
        return "light_green"
    if max_drawdown_usd < 250.0:
        return "yellow"
    if max_drawdown_usd < 300.0:
        return "light_red"
    return "dark_red"

def export_results_xlsx(path: str, all_results: dict, sort_by: str, sort_order: str, min_bets: int = 0):
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        log.error(
            "openpyxl isn't installed, so %s can't be written. Install it with "
            "'pip install openpyxl --break-system-packages', or use --export-csv instead "
            "(no extra packages needed, and CSV opens fine in Excel too).", path,
        )
        return

    FILLS = {
        "dark_green": PatternFill(start_color="375623", end_color="375623", fill_type="solid"),
        "light_green": PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),
        "yellow": PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid"),
        "light_red": PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),
        "dark_red": PatternFill(start_color="C00000", end_color="C00000", fill_type="solid"),
        "low_sample": PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid"),
    }
    FONTS = {
        "dark_green": Font(color="FFFFFF"),
        "light_green": Font(color="006100"),
        "yellow": Font(color="9C6500"),
        "light_red": Font(color="9C0006"),
        "dark_red": Font(color="FFFFFF"),
        "low_sample": Font(color="595959"),
    }

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for strategy, results in all_results.items():
        ranked = rank_results(results, sort_by, sort_order)

        param_columns = _varying_param_columns([params for params, _ in ranked])
        headers = param_columns + STAT_COLUMNS

        ws = wb.create_sheet(title=strategy[:31])
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)

        for params, r in ranked:
            stats = result_stat_row(r)
            row = [params.get(tuple(col.split(".")), "") for col in param_columns]
            row.extend(stats[c] for c in STAT_COLUMNS)
            ws.append(row)

            bucket = "low_sample" if r.bets < min_bets else _drawdown_color_bucket(r.max_drawdown_cents / 100.0)
            for cell in ws[ws.max_row]:
                cell.fill = FILLS[bucket]
                cell.font = FONTS[bucket]

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for i, col in enumerate(headers, start=1):
            ws.column_dimensions[get_column_letter(i)].width = max(10, min(40, len(col) + 2))

    wb.save(path)
    log.info("Wrote %s combo(s) across %s sheet(s) to %s", sum(len(v) for v in all_results.values()), len(all_results), path)

# ---------------------------------------------------------------------------
# Output
# ---------------------------------------------------------------------------

def _fmt_eta(seconds: float) -> str:
    seconds = int(max(0, seconds))
    h, rem = divmod(seconds, 3600)
    m, s = divmod(rem, 60)
    return f"{h}:{m:02d}:{s:02d}" if h else f"{m}:{s:02d}"

class ProgressBar:
    def __init__(self, total: int, prefix: str, enabled: bool = True, is_tty: bool = True, width: int = 30):
        self.total = max(1, total)
        self.prefix = prefix
        self.enabled = enabled
        self.is_tty = is_tty
        self.width = width
        self.start = time.monotonic()
        self._last_logged_pct = -1

    def update(self, current: int):
        if not self.enabled:
            return
        frac = min(1.0, current / self.total)
        elapsed = time.monotonic() - self.start
        eta = (elapsed / current * (self.total - current)) if current > 0 else 0.0
        if self.is_tty:
            filled = int(self.width * frac)
            bar = "#" * filled + "-" * (self.width - filled)
            sys.stdout.write(f"\r{self.prefix} [{bar}] {current}/{self.total} ({frac * 100:.0f}%) ETA {_fmt_eta(eta)}   ")
            sys.stdout.flush()
        else:
            pct = int(frac * 100)
            if pct != self._last_logged_pct and (pct % 10 == 0 or current == self.total):
                self._last_logged_pct = pct
                log.info("%s: %d/%d (%d%%) ETA %s", self.prefix, current, self.total, pct, _fmt_eta(eta))

    def finish(self):
        if self.enabled and self.is_tty:
            sys.stdout.write("\n")
            sys.stdout.flush()

def print_results_table(results: list, top_n: int, min_bets: int,
                         label_fields: Optional[list] = None, sort_by: str = "pnl", sort_order: str = "desc"):
    ranked = rank_results(results, sort_by, sort_order)[:top_n]

    rows = [(build_label(params, label_fields), r) for params, r in ranked]
    label_width = max([len(label) for label, _ in rows] + [len("Params")])

    print(f"(sorted by {sort_by}, {sort_order})")
    header = (
        f"{'Params':<{label_width}} {'Bets':>6} {'Wins':>6} {'Losses':>7} {'Skipped':>8} "
        f"{'WinRate':>8} {'MaxDrawdown':>12} {'PnL':>10}"
    )
    print(header)
    print("-" * len(header))
    for combo_label, r in rows:
        win_rate_str = f"{r.win_rate:.1f}%" if r.win_rate is not None else "n/a"
        pnl_str = f"${r.realized_pnl_cents / 100:+.2f}"
        drawdown_str = f"${r.max_drawdown_cents / 100:.2f}"
        flag = "  (LOW SAMPLE)" if r.bets < min_bets else ""
        print(
            f"{combo_label:<{label_width}} {r.bets:>6} {r.wins:>6} {r.losses:>7} {r.skipped:>8} "
            f"{win_rate_str:>8} {drawdown_str:>12} {pnl_str:>10}{flag}"
        )
    print()
    print(
        "PnL here is REAL, not assumed - it's computed from the actual up/down prices your bot recorded "
        "at the moment each simulated bet (and any hedge) would have qualified. MaxDrawdown is the worst "
        "peak of (unrecovered loss + cost of the bet just placed) reached in any single contiguous run of "
        "your recorded data - the same figure bot.py itself tracks live via update_max_drawdown(). It's "
        "cash at risk at the worst moment, not a running P&L low point, and it resets at gaps in your "
        "recordings since sizing state resets there too. 'Skipped' means either no qualifying price/signal "
        "ever appeared in the entry window, or that window's settlement couldn't be determined from your "
        "recorded data (e.g. a gap right at the end of a session). Rows flagged LOW SAMPLE had too few bets "
        "to trust the win rate / PnL / drawdown as more than noise."
    )

# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser(
        description="Backtest KXBTC15M bot strategies against recorded live-tick JSONL data (no Kalshi API calls)."
    )
    parser.add_argument("--backtest-config", default="backtest_ticks_config.yaml")
    parser.add_argument("--data-dir", default=None, help="Override data_dir from the backtest config")
    parser.add_argument("--top", type=int, default=25, help="Show only the top N rows per strategy")
    parser.add_argument("--min-bets", type=int, default=20, help="Flag rows with fewer bets than this as low-sample")
    parser.add_argument("--max-combos", type=int, default=200, help="Safety cap on grid size per strategy")
    parser.add_argument(
        "--label-fields", default=None,
        help="Comma-separated swept param names to show in the Params column, e.g. "
             "'spot_lean.threshold_pct,sizing.mode'. Defaults to the backtest config's "
             "label_fields, or every swept param if that's also unset.",
    )
    parser.add_argument(
        "--sort-by", default=None, choices=list(SORT_KEYS.keys()),
        help="Column to sort results by. Defaults to the backtest config's sort_by, or 'pnl'.",
    )
    parser.add_argument(
        "--sort-order", default=None, choices=["asc", "desc"],
        help="Sort direction. Defaults to the backtest config's sort_order, or 'desc'.",
    )
    parser.add_argument(
        "--export-csv", default=None,
        help="Write EVERY combo tested (not just --top) to this CSV file, one row per combo, "
             "each swept param as its own column. Opens fine in Excel/Sheets, no extra packages needed.",
    )
    parser.add_argument(
        "--export-xlsx", default=None,
        help="Write EVERY combo tested to this .xlsx file, one sheet per strategy. Requires "
             "'pip install openpyxl'.",
    )
    parser.add_argument(
        "--quiet", action="store_true",
        help="Skip printing the results table(s) to the console - useful when you're only after "
             "--export-csv/--export-xlsx. Progress logging (loading data, running combos) still shows.",
    )
    parser.add_argument(
        "--no-progress", action="store_true",
        help="Disable the live progress bar while combos run.",
    )
    args = parser.parse_args()

    logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(name)s: %(message)s")
    logging.getLogger("bot").setLevel(logging.WARNING)
    logging.getLogger("state").setLevel(logging.ERROR)

    bt_cfg = load_yaml(args.backtest_config)
    data_dir = args.data_dir or bt_cfg.get("data_dir", "./data")
    file_pattern = bt_cfg.get("file_pattern", "*.jsonl")
    base_cfg = load_yaml(bt_cfg.get("base_config", "./config.yaml"))
    strategies = bt_cfg.get("strategies_to_test", ["spot_lean"])
    grid_cfg = bt_cfg.get("grid", {})
    common_grid = grid_cfg.get("common", {})

    label_fields_raw = args.label_fields if args.label_fields is not None else bt_cfg.get("label_fields")
    label_fields = [f.strip() for f in label_fields_raw.split(",")] if isinstance(label_fields_raw, str) else label_fields_raw
    sort_by = args.sort_by or bt_cfg.get("sort_by", "pnl")
    sort_order = args.sort_order or bt_cfg.get("sort_order", "desc")
    if sort_by not in SORT_KEYS:
        log.warning("Unknown sort_by=%s, falling back to 'pnl'. Valid options: %s", sort_by, ", ".join(SORT_KEYS))
        sort_by = "pnl"

    log.info("Loading tick data from %s (pattern: %s)...", data_dir, file_pattern)
    windows = load_all_windows(data_dir, file_pattern)
    if not windows:
        log.error("No usable JSONL tick files found in %s - nothing to backtest.", data_dir)
        sys.exit(1)

    segments = contiguous_segments(windows)
    total_windows = sum(len(s) for s in segments)
    log.info("Loaded %s windows across %s contiguous segment(s) (gaps reset sizing state, same as backtest.py).",
              total_windows, len(segments))
    if total_windows < 50:
        log.warning("Only %s windows of tick data available - results below will have wide uncertainty.", total_windows)

    results_by_segment = []
    for seg in segments:
        seg_results = []
        for i, w in enumerate(seg):
            next_w = seg[i + 1] if i + 1 < len(seg) else None
            seg_results.append(determine_result(w, next_w))
        results_by_segment.append(seg_results)

    spot_index = build_spot_index(windows)

    all_results = {}
    for strategy in strategies:
        strat_grid = grid_cfg.get(strategy, {})
        combos = list(iter_grid_combos(base_cfg, common_grid, strat_grid))
        if len(combos) > args.max_combos:
            log.warning("%s: grid has %s combos, capping at %s (raise --max-combos to test more).",
                        strategy, len(combos), args.max_combos)
            combos = combos[:args.max_combos]

        log.info("Running %s combo(s) for strategy=%s ...", len(combos), strategy)
        results = []
        bar = ProgressBar(len(combos), prefix=strategy, enabled=not args.no_progress, is_tty=sys.stdout.isatty())
        for i, (params, cfg) in enumerate(combos, 1):
            cfg.setdefault("strategy", {})["mode"] = strategy
            agg = run_strategy_over_segments(strategy, segments, results_by_segment, spot_index, cfg)
            results.append((params, agg))
            bar.update(i)
        bar.finish()

        if not args.quiet:
            print()
            print(f"===== {strategy.upper()} - tick-data backtest ({total_windows} windows, {len(segments)} segment(s)) =====")
            print_results_table(results, args.top, args.min_bets, label_fields, sort_by, sort_order)
        all_results[strategy] = results

    if args.export_csv:
        export_results_csv(args.export_csv, all_results, sort_by, sort_order)
    if args.export_xlsx:
        export_results_xlsx(args.export_xlsx, all_results, sort_by, sort_order, args.min_bets)

if __name__ == "__main__":
    main()